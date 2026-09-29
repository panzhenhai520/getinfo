import io
import os
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
import openpyxl

import intel_reports
from intel_reports import (
    MAX_REPORT_CONVERT_BYTES,
    IntelReportService,
)
from sqlite_database import SQLiteDatabase


def build_docx_bytes(title="2026 网络安全行业研究报告"):
    """构造真实 .docx 字节：标题 + 段落 + 表格。"""
    doc = Document()
    doc.add_heading(title, level=1)
    doc.add_paragraph("本报告聚焦网络安全与工控安全领域的年度态势。")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "指标"
    table.cell(0, 1).text = "数值"
    table.cell(1, 0).text = "漏洞数量"
    table.cell(1, 1).text = "1286"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_xlsx_bytes():
    """构造真实 .xlsx 字节：含表格数据。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "威胁态势"
    ws.append(["季度", "工控漏洞数"])
    ws.append(["Q1", "320"])
    ws.append(["Q2", "415"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class IntelReportsMarkitdownTests(unittest.TestCase):
    """阶段2：报告文件（Word/Excel/PPT）markitdown 内联阅读"""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        os.environ["INTEL_REPORT_STORAGE_DIR"] = os.path.join(self.tmp.name, "reports")
        self.db = SQLiteDatabase(os.path.join(self.tmp.name, "reports.sqlite3"))
        self.assertTrue(self.db.connect())
        self.assertTrue(self.db.create_tables())
        self.service = IntelReportService(self.db)

    def tearDown(self):
        self.db.disconnect()
        self.tmp.cleanup()
        os.environ.pop("INTEL_REPORT_STORAGE_DIR", None)

    def _ingest_patches(self):
        import contextlib
        stack = contextlib.ExitStack()
        stack.enter_context(
            patch.object(intel_reports.intel_llm_client, "extract_report_items", return_value=[])
        )
        stack.enter_context(
            patch.object(intel_reports, "resolve_industry_ragflow_kb_id", return_value="")
        )
        # 主题匹配依赖具体行业包配置，单测按"通过"处理（主题匹配逻辑另有测试覆盖）
        stack.enter_context(
            patch.object(intel_reports, "report_matches_industry_topic", return_value=True)
        )
        return stack

    def test_office_format_hint(self):
        hint = IntelReportService._office_format_hint
        self.assertEqual(hint("", "报告.docx"), "docx")
        self.assertEqual(hint("", "数据.xlsx"), "xlsx")
        self.assertEqual(hint("", "演示.PPTX"), "pptx")
        self.assertEqual(
            hint("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", ""),
            "xlsx",
        )
        self.assertEqual(hint("", "bad.zip"), "")
        self.assertEqual(hint("application/pdf", ""), "")

    def test_unsupported_extension_rejected_on_upload(self):
        with self._ingest_patches():
            with self.assertRaises(ValueError):
                self.service.ingest_upload(b"not-a-doc", "file.zip", "family_office")

    def test_ingest_docx_creates_markdown_asset_with_table(self):
        with self._ingest_patches():
            result = self.service.ingest_upload(
                build_docx_bytes(), "安全报告.docx", "family_office", title="网络安全行业研究报告"
            )
        self.assertGreater(result["report_id"], 0)
        report, path, mimetype = self.service.report_asset(result["report_id"], "markdown")
        self.assertIn("text/markdown", mimetype)
        text = path.read_text(encoding="utf-8")
        self.assertIn("网络安全行业研究报告", text)
        self.assertIn("漏洞数量", text)
        self.assertIn("|", text)  # 表格结构保留

        # 原始文件下载的 MIME 正确（Word）
        _report, _raw_path, original_mime = self.service.report_asset(result["report_id"], "original")
        self.assertIn("wordprocessingml", original_mime)

    def test_ingest_xlsx_creates_markdown_asset_with_table(self):
        with self._ingest_patches():
            result = self.service.ingest_upload(
                build_xlsx_bytes(), "威胁态势.xlsx", "family_office", title="工控安全威胁态势统计"
            )
        _report, path, _mimetype = self.service.report_asset(result["report_id"], "markdown")
        text = path.read_text(encoding="utf-8")
        self.assertIn("工控安全威胁态势统计", text)
        self.assertIn("工控漏洞数", text)
        self.assertIn("|", text)

    def test_metadata_records_original_format(self):
        import json
        from db_connection import is_postgres_connection
        with self._ingest_patches():
            result = self.service.ingest_upload(
                build_docx_bytes(), "安全报告.docx", "family_office", title="网络安全行业研究报告"
            )
        self.db._ensure_connection()
        with self.db.lock:
            row = self.db.connection.execute(
                "SELECT metadata_json FROM intel_reports WHERE id=?", (result["report_id"],)
            ).fetchone()
        meta = json.loads(row["metadata_json"])
        self.assertEqual(meta.get("original_format"), "docx")

    def test_too_large_office_file_rejected_for_conversion(self):
        docx = build_docx_bytes()
        with patch.object(intel_reports, "MAX_REPORT_CONVERT_BYTES", 10):
            # 写入临时文件再转换（>10 字节即触发上限）
            import pathlib
            path = pathlib.Path(self.tmp.name) / "big.docx"
            path.write_bytes(docx)
            with self.assertRaises(ValueError) as ctx:
                self.service._office_to_markdown(path, "docx")
        self.assertIn("文件过大", str(ctx.exception))

    def test_missing_file_rejected(self):
        import pathlib
        with self.assertRaises(ValueError):
            self.service._office_to_markdown(pathlib.Path(self.tmp.name) / "nope.docx", "docx")


if __name__ == "__main__":
    unittest.main()
