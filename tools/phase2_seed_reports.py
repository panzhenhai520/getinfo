# -*- coding: utf-8 -*-
"""阶段2 验收准备：本地为 bolean 行业包登记一份 Word + 一份 Excel 报告。"""
import io
import sys

sys.path.insert(0, ".")

from docx import Document
import openpyxl

from intel_reports import intel_report_service

PACK = "bolean_security_compute"


def docx_bytes():
    doc = Document()
    doc.add_heading("网络安全行业研究报告（内联阅读验收）", level=1)
    doc.add_paragraph("本报告聚焦网络安全与工控安全年度态势，供报告页内联阅读验收使用。")
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "指标"
    table.cell(0, 1).text = "数值"
    table.cell(1, 0).text = "高危漏洞"
    table.cell(1, 1).text = "1286"
    table.cell(2, 0).text = "受攻击行业"
    table.cell(2, 1).text = "制造业"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def xlsx_bytes():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "威胁态势"
    ws.append(["季度", "工控漏洞数"])
    ws.append(["Q1", "320"])
    ws.append(["Q2", "415"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


r1 = intel_report_service.ingest_upload(
    docx_bytes(), "网络安全行业研究报告.docx", PACK, title="网络安全行业研究报告（内联阅读验收）"
)
print("docx report:", r1.get("report_id"), r1.get("status"))
r2 = intel_report_service.ingest_upload(
    xlsx_bytes(), "工控安全威胁态势.xlsx", PACK, title="工控安全威胁态势统计（内联阅读验收）"
)
print("xlsx report:", r2.get("report_id"), r2.get("status"))
