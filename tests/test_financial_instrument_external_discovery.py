import io
import json
import tempfile
import unittest
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import parse_qs, urlsplit

from openpyxl import Workbook

from financial_instrument_discovery import FinancialInstrumentDiscoveryService
from financial_instrument_sources import (
    EASTMONEY_HK_PROFILE_URL,
    EASTMONEY_HK_QUOTE_IDENTITY_FALLBACK_URL,
    EASTMONEY_HK_QUOTE_IDENTITY_URL,
    HKEX_SECURITIES_URL,
    HKEX_SECURITIES_ZH_URL,
    NASDAQ_LISTED_URL,
    NASDAQ_OTHER_URL,
    SSE_STOCK_LIST_URL,
    SZSE_STOCK_LIST_URL,
    AKShareAInstrumentDiscoverySource,
    AKShareHKInstrumentDiscoverySource,
    AlphaVantageInstrumentDiscoverySource,
    HKEXInstrumentDiscoverySource,
    NasdaqTraderInstrumentDiscoverySource,
    SSEInstrumentDiscoverySource,
    SerpAPIInstrumentHintSource,
    SZSEInstrumentDiscoverySource,
    YahooInstrumentDiscoverySource,
)
from financial_instruments import InstrumentRegistry
from sqlite_database import SQLiteDatabase


UTC = timezone.utc
NOW = datetime(2026, 8, 4, 12, 0, tzinfo=UTC)
SETTINGS = {
    "FINANCIAL_INTELLIGENCE_ENABLED": True,
    "FINANCIAL_INSTRUMENT_DISCOVERY_ENABLED": True,
    "FINANCIAL_INSTRUMENT_AUTO_PROMOTION_ENABLED": True,
    "FINANCIAL_ROLLOUT_STAGE": "simulation_backtest",
    "FINANCIAL_LICENSE_ENVIRONMENT": "development",
    "ALPHA_VANTAGE_ENABLED": True,
    "ALPHA_VANTAGE_API_KEY": "fixture-key",
    "AKSHARE_CN_ENABLED": True,
    "YAHOO_FINANCE_ENABLED": True,
}


class _Result:
    def __init__(self, content, url):
        self.content = content if isinstance(content, bytes) else content.encode("utf-8")
        self.url = url

    @property
    def text(self):
        return self.content.decode("utf-8")


class _YahooSearchResult:
    def __init__(self, query):
        fixtures = {
            "NUVB": {
                "symbol": "NUVB",
                "shortname": "Nuvation Bio Inc.",
                "longname": "Nuvation Bio Inc.",
                "quoteType": "EQUITY",
                "exchange": "NMS",
                "exchDisp": "NASDAQ",
            }
        }
        self.quotes = [fixtures[query]] if query in fixtures else []


class _YahooSDK:
    @staticmethod
    def Search(query, **kwargs):
        del kwargs
        return _YahooSearchResult(query)


def _hkex_workbook():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["List of Securities"])
    sheet.append(["Updated as at 04/08/2026"])
    sheet.append(
        [
            "Stock Code",
            "Name of Securities",
            "Category",
            "Sub-Category",
            "ISIN",
            "Trading Currency",
        ]
    )
    sheet.append(
        [
            "09969",
            "INNOCARE",
            "Equity",
            "Equity Securities (Main Board)",
            "KYG4783B1032",
            "HKD",
        ]
    )
    sheet.append(
        [
            "03119",
            "GX ASIA SEMICON",
            "Exchange Traded Products",
            "Exchange Traded Funds",
            "HK0000756236",
            "HKD",
        ]
    )
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _hkex_zh_workbook():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["證券名單"])
    sheet.append(["截 至 04/08/2026"])
    sheet.append(["股份代號", "股份名稱", "分類", "次分類", "國際證券號碼 (ISIN)", "交易貨幣"])
    sheet.append(["09969", "諾誠健華", "股本", "股本證券(主板)", "KYG4783B1032", "HKD"])
    sheet.append(["03119", "GX亞洲半導體", "交易所買賣產品", "交易所買賣基金", "HK0000756236", "HKD"])
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _szse_workbook():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["深圳证券交易所股票列表"])
    sheet.append(
        [
            "板块",
            "A股代码",
            "A股简称",
            "A股上市日期",
            "A股总股本",
            "A股流通股本",
            "所属行业",
        ]
    )
    sheet.append(
        ["创业板", "300750", "宁德时代", "2018-06-11", 0, 0, "电气机械和器材制造业"]
    )
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


class _HTTP:
    def __init__(self):
        self.calls = []
        self.alpha = {
            "NUVB": {
                "1. symbol": "NUVB",
                "2. name": "Nuvation Bio Inc.",
                "3. type": "Equity",
                "4. region": "United States",
                "8. currency": "USD",
            },
            "9969": {
                "1. symbol": "9969.HKG",
                "2. name": "InnoCare Pharma Limited",
                "3. type": "Equity",
                "4. region": "Hong Kong",
                "8. currency": "HKD",
            },
            "3119": {
                "1. symbol": "3119.HKG",
                "2. name": "Global X Asia Semiconductor ETF",
                "3. type": "ETF",
                "4. region": "Hong Kong",
                "8. currency": "HKD",
            },
            "Nuvation Bio Inc": {
                "1. symbol": "NUVB",
                "2. name": "Nuvation Bio Inc.",
                "3. type": "Equity",
                "4. region": "United States",
                "8. currency": "USD",
            },
        }

    def get(self, url, *, headers=None):
        del headers
        self.calls.append(url)
        if url == NASDAQ_LISTED_URL:
            return _Result(
                "Symbol|Security Name|Market Category|Test Issue|Financial Status|Round Lot Size|ETF|NextShares\n"
                "NUVB|Nuvation Bio Inc. - Class A Common Stock|Q|N|N|100|N|N\n"
                "USXF|US Example Fund ETF|G|N|N|100|Y|N\n"
                "File Creation Time: 0804202621:31|||||||\n",
                url,
            )
        if url == NASDAQ_OTHER_URL:
            return _Result(
                "ACT Symbol|Security Name|Exchange|CQS Symbol|ETF|Round Lot Size|Test Issue|NASDAQ Symbol\n"
                "ZZZZ|Example NYSE Equity|N|ZZZZ|N|100|N|ZZZZ\n"
                "File Creation Time: 0804202621:31|||||||\n",
                url,
            )
        if url == HKEX_SECURITIES_URL:
            return _Result(_hkex_workbook(), url)
        if url == HKEX_SECURITIES_ZH_URL:
            return _Result(_hkex_zh_workbook(), url)
        if url.startswith(SSE_STOCK_LIST_URL):
            stock_type = parse_qs(urlsplit(url).query).get("STOCK_TYPE", [""])[0]
            rows = (
                [
                    {
                        "A_STOCK_CODE": "601919",
                        "SEC_NAME_CN": "中远海控",
                        "SEC_NAME_FULL": "中远海运控股股份有限公司",
                        "COMPANY_ABBR": "中远海控",
                        "FULL_NAME": "中远海运控股股份有限公司",
                        "LIST_DATE": "20070626",
                    }
                ]
                if stock_type == "1"
                else []
            )
            return _Result(json.dumps({"result": rows}), url)
        if url.startswith(SZSE_STOCK_LIST_URL):
            return _Result(_szse_workbook(), url)
        if url.startswith(EASTMONEY_HK_PROFILE_URL):
            query = parse_qs(urlsplit(url).query)
            code = query.get("filter", [""])[0].split('"')[1].split(".", 1)[0]
            fixtures = {
                "09969": {
                    "SECUCODE": "09969.HK",
                    "SECURITY_CODE": "09969",
                    "SECURITY_NAME_ABBR": "诺诚健华",
                    "SECURITY_TYPE": "普通股",
                    "LISTING_DATE": "2020-03-23",
                    "ISIN_CODE": "KYG4783B1032",
                    "BOARD": "主板",
                    "TRADE_MARKET": "香港交易所",
                },
                "03119": {
                    "SECUCODE": "03119.HK",
                    "SECURITY_CODE": "03119",
                    "SECURITY_NAME_ABBR": "GX亚洲半导体",
                    "SECURITY_TYPE": "ETF",
                    "LISTING_DATE": "2021-08-24",
                    "ISIN_CODE": "HK0000756236",
                    "BOARD": "ETF",
                    "TRADE_MARKET": "香港交易所",
                },
            }
            rows = [fixtures[code]] if code in fixtures else []
            if code == "03119":
                rows = []
            return _Result(json.dumps({"result": {"data": rows}}), url)
        if url.startswith(
            (
                EASTMONEY_HK_QUOTE_IDENTITY_URL,
                EASTMONEY_HK_QUOTE_IDENTITY_FALLBACK_URL,
            )
        ):
            secid = parse_qs(urlsplit(url).query).get("secid", [""])[0]
            fixtures = {
                "116.03119": {
                    "f57": "03119",
                    "f58": "GX亚洲半导体",
                    "f107": 116,
                },
                "1.601919": {"f57": "601919", "f58": "中远海控", "f107": 1},
                "0.300750": {"f57": "300750", "f58": "宁德时代", "f107": 0},
            }
            return _Result(json.dumps({"data": fixtures.get(secid)}), url)
        if "alphavantage.co" in url:
            keyword = parse_qs(urlsplit(url).query).get("keywords", [""])[0]
            payload = {"bestMatches": [self.alpha[keyword]] if keyword in self.alpha else []}
            return _Result(json.dumps(payload), url)
        raise AssertionError(f"unexpected URL: {url}")


class _PrimaryQuoteFailsHTTP(_HTTP):
    def get(self, url, *, headers=None):
        if url.startswith(EASTMONEY_HK_QUOTE_IDENTITY_URL):
            self.calls.append(url)
            raise RuntimeError("primary quote identity endpoint unavailable")
        return super().get(url, headers=headers)


class _WrongCNMarketMarkerHTTP(_HTTP):
    def get(self, url, *, headers=None):
        if url.startswith(
            (EASTMONEY_HK_QUOTE_IDENTITY_URL, EASTMONEY_HK_QUOTE_IDENTITY_FALLBACK_URL)
        ) and parse_qs(urlsplit(url).query).get("secid", [""])[0] == "1.601919":
            self.calls.append(url)
            return _Result(
                json.dumps(
                    {"data": {"f57": "601919", "f58": "中远海控", "f107": 0}}
                ),
                url,
            )
        return super().get(url, headers=headers)


class _SearchClient:
    def __init__(self, results):
        self.results = results
        self.calls = []

    def search(self, query, *, recency_days=None):
        self.calls.append((query, recency_days))
        return list(self.results)


class FinancialInstrumentExternalDiscoveryTest(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.database = SQLiteDatabase(
            str(Path(self.temp_dir.name) / "external-discovery.sqlite3")
        )
        self.assertTrue(self.database.connect())
        self.assertTrue(self.database.create_tables())
        self.http = _HTTP()

    def tearDown(self):
        self.database.disconnect()
        self.temp_dir.cleanup()

    def _structured_sources(self):
        return (
            NasdaqTraderInstrumentDiscoverySource(http_client=self.http),
            SSEInstrumentDiscoverySource(http_client=self.http),
            SZSEInstrumentDiscoverySource(http_client=self.http),
            HKEXInstrumentDiscoverySource(http_client=self.http),
            AKShareAInstrumentDiscoverySource(
                settings=SETTINGS, http_client=self.http
            ),
            AKShareHKInstrumentDiscoverySource(
                settings=SETTINGS, http_client=self.http
            ),
            AlphaVantageInstrumentDiscoverySource(
                settings=SETTINGS, http_client=self.http
            ),
        )

    def _discover(self, question, *, sources=None):
        service = FinancialInstrumentDiscoveryService(
            self.database.connection,
            sources=sources or self._structured_sources(),
            settings=SETTINGS,
        )
        return service.discover_and_promote(
            question, requested_at=NOW, request_id="external-fixture"
        )

    def test_us_equity_uses_exchange_directory_and_approved_provider_mapping(self):
        result = self._discover("NUVB 美股最新情况")

        self.assertEqual(result["status"], "promoted")
        self.assertEqual(result["promoted_target"]["canonical_symbol"], "NUVB.US")
        self.assertEqual(result["promoted_target"]["exchange"], "XNAS")
        self.assertEqual(
            result["promoted_target"]["provider_mappings"]["alpha_vantage"],
            "NUVB",
        )
        self.assertEqual(
            result["promoted_target"]["provider_mappings"]["yahoo"],
            "NUVB",
        )
        self.assertFalse(
            result["promoted_target"]["metadata"]["provider_mapping_derivation"][
                "counts_as_identity_corroboration"
            ]
        )

    def test_hkex_official_bilingual_names_are_news_aliases(self):
        source = HKEXInstrumentDiscoverySource(http_client=self.http)

        assertions = source.search(
            "諾誠健華", requested_at=NOW, request_id="hkex-bilingual"
        )

        self.assertEqual(len(assertions), 1)
        self.assertEqual(assertions[0]["canonical_symbol"], "9969.HK")
        self.assertEqual(assertions[0]["display_name"], "INNOCARE")
        self.assertIn("諾誠健華", assertions[0]["aliases"])
        self.assertEqual(
            assertions[0]["metadata"]["official_names"],
            {"en_short": "INNOCARE", "zh_short": "諾誠健華"},
        )

    def test_us_equity_can_use_yahoo_exact_search_when_alpha_is_unavailable(self):
        result = self._discover(
            "NUVB.US 股票最新情况",
            sources=(
                NasdaqTraderInstrumentDiscoverySource(http_client=self.http),
                YahooInstrumentDiscoverySource(settings=SETTINGS, sdk=_YahooSDK()),
            ),
        )

        self.assertEqual(result["status"], "promoted")
        self.assertEqual(result["promoted_target"]["canonical_symbol"], "NUVB.US")
        self.assertEqual(
            result["promoted_target"]["provider_mappings"]["yahoo"], "NUVB"
        )
        self.assertIn(
            "yahoo_finance_symbol_search",
            result["promoted_target"]["metadata"]["discovery_source_keys"],
        )

    def test_us_name_uses_the_same_structured_admission_gate(self):
        result = self._discover("Nuvation Bio Inc. 股票怎么样")

        self.assertEqual(result["status"], "promoted")
        self.assertEqual(result["promoted_target"]["canonical_symbol"], "NUVB.US")
        self.assertEqual(
            result["promoted_target"]["provider_mappings"]["alpha_vantage"],
            "NUVB",
        )

    def test_cn_exchange_lists_and_provider_mapping_support_code_and_name(self):
        cases = (
            ("601919.SH 股票最新信息", "601919.SH", "XSHG", "601919"),
            ("中远海控股票怎么样", "601919.SH", "XSHG", "601919"),
            ("300750.SZ 最新股价", "300750.SZ", "XSHE", "300750"),
            ("宁德时代股票怎么样", "300750.SZ", "XSHE", "300750"),
        )
        for question, symbol, exchange, provider_symbol in cases:
            with self.subTest(question=question):
                result = self._discover(question)
                self.assertEqual(result["status"], "promoted")
                self.assertEqual(
                    result["promoted_target"]["canonical_symbol"], symbol
                )
                self.assertEqual(result["promoted_target"]["exchange"], exchange)
                self.assertRegex(
                    result["promoted_target"]["listed_at"], r"^\d{4}-\d{2}-\d{2}$"
                )
                self.assertEqual(
                    result["promoted_target"]["provider_mappings"]["akshare_cn"],
                    provider_symbol,
                )
                suffix = symbol.rsplit(".", 1)[1]
                self.assertEqual(
                    result["promoted_target"]["provider_mappings"]["tushare_cn"],
                    symbol,
                )
                self.assertEqual(
                    result["promoted_target"]["provider_mappings"]["yahoo"],
                    f"{provider_symbol}.{'SS' if suffix == 'SH' else 'SZ'}",
                )
        self.assertEqual(
            self.database.connection.execute(
                "SELECT COUNT(*) FROM financial_instruments "
                "WHERE canonical_symbol IN ('601919.SH', '300750.SZ')"
            ).fetchone()[0],
            2,
        )

    def test_cn_name_requeries_provider_with_exchange_verified_symbol(self):
        result = self._discover("中远海控股票怎么样")

        self.assertEqual(result["status"], "promoted")
        self.assertIn(
            "akshare_eastmoney_a_share_quote_identity:canonical-1",
            result["attempted_sources"],
        )

    def test_cn_quote_identity_alone_or_wrong_market_marker_never_promotes(self):
        quote_only = self._discover(
            "601919.SH 股票最新信息",
            sources=(
                AKShareAInstrumentDiscoverySource(
                    settings=SETTINGS, http_client=self.http
                ),
            ),
        )
        self.assertEqual(quote_only["status"], "verification_required")
        self.assertIsNone(
            InstrumentRegistry(self.database.connection).get_by_canonical_symbol(
                "601919.SH"
            )
        )

        self.http = _WrongCNMarketMarkerHTTP()
        wrong_market = self._discover(
            "601919.SH 股票最新信息",
            sources=(
                SSEInstrumentDiscoverySource(http_client=self.http),
                AKShareAInstrumentDiscoverySource(
                    settings=SETTINGS, http_client=self.http
                ),
            ),
        )
        self.assertEqual(wrong_market["status"], "rejected")
        self.assertIn(
            "independent_corroboration_required", wrong_market["reason_codes"]
        )
        self.assertIsNone(
            InstrumentRegistry(self.database.connection).get_by_canonical_symbol(
                "601919.SH"
            )
        )

    def test_bare_cn_code_remains_unresolved_and_is_never_registered(self):
        result = self._discover("601919 股票怎么样")

        self.assertEqual(result["status"], "not_found")
        self.assertIsNone(
            InstrumentRegistry(self.database.connection).get_by_canonical_symbol(
                "601919.SH"
            )
        )

    def test_hk_equity_and_etf_keep_distinct_asset_types(self):
        equity = self._discover("港股 9969.HK 最新情况")
        fund = self._discover("3119.HK 基金走势")

        self.assertEqual(equity["promoted_target"]["canonical_symbol"], "9969.HK")
        self.assertEqual(equity["promoted_target"]["asset_type"], "equity")
        self.assertEqual(fund["promoted_target"]["canonical_symbol"], "3119.HK")
        self.assertEqual(fund["promoted_target"]["asset_type"], "etf")

    def test_hk_name_is_resolved_by_exchange_then_rechecked_by_provider(self):
        result = self._discover("INNOCARE 股票怎么样")

        self.assertEqual(result["status"], "promoted")
        self.assertEqual(result["promoted_target"]["canonical_symbol"], "9969.HK")
        self.assertEqual(
            result["promoted_target"]["provider_mappings"]["akshare_cn"],
            "09969",
        )
        self.assertEqual(
            result["promoted_target"]["provider_mappings"]["alpha_vantage"],
            "9969.HKG",
        )
        self.assertEqual(
            result["promoted_target"]["provider_mappings"]["yahoo"],
            "9969.HK",
        )
        self.assertIn(
            "akshare_eastmoney_hk_security_profile:canonical-1",
            result["attempted_sources"],
        )
        alias = self.database.connection.execute(
            """
            SELECT alias_type, source_key, source_url, is_official
            FROM financial_instrument_aliases
            WHERE instrument_id=? AND alias='諾誠健華'
            """,
            (result["promoted_target"]["instrument_id"],),
        ).fetchone()
        self.assertEqual(
            tuple(alias),
            (
                "official_zh_short",
                "hkex_full_list_of_securities",
                HKEX_SECURITIES_URL,
                1,
            ),
        )

    def test_hk_etf_quote_identity_fallback_uses_exchange_type_and_keeps_gate(self):
        self.http = _PrimaryQuoteFailsHTTP()
        sources = (
            HKEXInstrumentDiscoverySource(http_client=self.http),
            AKShareHKInstrumentDiscoverySource(
                settings=SETTINGS, http_client=self.http
            ),
        )

        result = self._discover("3119.HK 基金走势", sources=sources)

        self.assertEqual(result["status"], "promoted")
        self.assertEqual(result["promoted_target"]["canonical_symbol"], "3119.HK")
        self.assertEqual(result["promoted_target"]["asset_type"], "etf")
        self.assertEqual(
            result["promoted_target"]["provider_mappings"]["akshare_cn"],
            "03119",
        )
        self.assertNotIn(
            "alpha_vantage",
            result["promoted_target"]["provider_mappings"],
        )
        self.assertEqual(
            result["promoted_target"]["provider_mappings"]["easyquotation"],
            "03119",
        )
        self.assertTrue(
            any(
                url.startswith(EASTMONEY_HK_QUOTE_IDENTITY_FALLBACK_URL)
                for url in self.http.calls
            )
        )

    def test_hk_etf_quote_identity_alone_cannot_cross_admission_gate(self):
        source = AKShareHKInstrumentDiscoverySource(
            settings=SETTINGS, http_client=self.http
        )

        result = self._discover("3119.HK 基金走势", sources=(source,))

        self.assertEqual(result["status"], "verification_required")
        self.assertIsNone(
            InstrumentRegistry(self.database.connection).get_by_canonical_symbol(
                "3119.HK"
            )
        )

    def test_search_hint_is_reverified_by_structured_sources_before_promotion(self):
        search = _SearchClient(
            [
                {
                    "title": "InnoCare Pharma Limited (HKEX: 9969)",
                    "summary": "港交所上市公司资料",
                    "url": "https://example.test/innocare",
                }
            ]
        )
        sources = (
            *self._structured_sources(),
            SerpAPIInstrumentHintSource(client=search),
        )

        result = self._discover("诺诚健华股票最新情况", sources=sources)

        self.assertEqual(result["status"], "promoted")
        self.assertEqual(result["promoted_target"]["canonical_symbol"], "9969.HK")
        self.assertEqual(result["search_hints"][0]["suggested_queries"], ["9969.HK"])
        self.assertTrue(
            any(item.endswith(":suggestion-1") for item in result["attempted_sources"])
        )

    def test_web_hint_without_structured_confirmation_never_promotes(self):
        search = _SearchClient(
            [
                {
                    "title": "Unknown Company (NASDAQ: QQQQQQ)",
                    "summary": "unverified",
                    "url": "https://example.test/unknown",
                }
            ]
        )
        sources = (
            *self._structured_sources(),
            SerpAPIInstrumentHintSource(client=search),
        )

        result = self._discover("完全陌生公司股票最新情况", sources=sources)

        self.assertEqual(result["status"], "not_found")
        self.assertIsNone(
            InstrumentRegistry(self.database.connection).get_by_canonical_symbol(
                "QQQQQQ.US"
            )
        )


if __name__ == "__main__":
    unittest.main()
