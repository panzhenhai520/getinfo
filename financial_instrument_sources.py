#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""结构化证券发现来源；搜索引擎只提供待复核代码，不提供准入证据。"""

from __future__ import annotations

import csv
import io
import json
import re
import threading
from datetime import datetime, timedelta, timezone
from typing import Mapping
from urllib.parse import urlencode

from financial_config import require_financial_capability
from financial_instruments import normalize_alias
from financial_providers.base import load_profile
from financial_source_license import require_provider_authorization
from intel_http import SafeHTTPClient
from serpapi_client import SerpAPIClient


UTC = timezone.utc
NASDAQ_LISTED_URL = (
    "https://www.nasdaqtrader.com/dynamic/SymDir/nasdaqlisted.txt"
)
NASDAQ_OTHER_URL = (
    "https://www.nasdaqtrader.com/dynamic/SymDir/otherlisted.txt"
)
HKEX_SECURITIES_URL = (
    "https://www.hkex.com.hk/eng/services/trading/securities/"
    "securitieslists/ListOfSecurities.xlsx"
)
HKEX_SECURITIES_ZH_URL = (
    "https://www.hkex.com.hk/chi/services/trading/securities/"
    "securitieslists/ListOfSecurities_c.xlsx"
)
SSE_STOCK_LIST_URL = "https://query.sse.com.cn/sseQuery/commonQuery.do"
SSE_STOCK_LIST_PAGE_URL = "https://www.sse.com.cn/assortment/stock/list/share/"
SZSE_STOCK_LIST_URL = "https://www.szse.cn/api/report/ShowReport"
SZSE_STOCK_LIST_PAGE_URL = "https://www.szse.cn/market/product/stock/list/index.html"
EASTMONEY_HK_PROFILE_URL = (
    "https://datacenter.eastmoney.com/securities/api/data/v1/get"
)
EASTMONEY_HK_QUOTE_IDENTITY_URL = (
    "https://push2.eastmoney.com/api/qt/stock/get"
)
EASTMONEY_HK_QUOTE_IDENTITY_FALLBACK_URL = (
    "https://push2delay.eastmoney.com/api/qt/stock/get"
)
_ALPHA_PROFILE = load_profile("alpha_vantage")
_AKSHARE_PROFILE = load_profile("akshare_cn")
_YAHOO_PROFILE = load_profile("yahoo")
_EXCHANGE_CODES = {
    "N": "XNYS",
    "A": "XASE",
    "P": "ARCX",
    "Z": "BATS",
    "V": "IEXG",
}
_UNSUPPORTED_SECURITY = re.compile(
    r"\b(?:warrant|rights?|units?|notes?|bonds?|debentures?)\b", re.I
)
_HK_SYMBOL = re.compile(r"^(?P<code>\d{1,5})(?:\.(?:HK|HKG))?$", re.I)
_US_SYMBOL = re.compile(r"^(?P<code>[A-Z][A-Z0-9.-]{0,14}?)(?:\.US)?$", re.I)
_CN_SYMBOL = re.compile(r"^(?P<code>\d{6})\.(?P<venue>SH|SZ)$", re.I)
_SUGGESTION_PATTERNS = (
    re.compile(r"\b(?P<code>\d{1,5})\s*\.(?:HK|HKG)\b", re.I),
    re.compile(r"\b(?:HKEX|SEHK|港交所)\s*[:：]\s*(?P<code>\d{1,5})\b", re.I),
    re.compile(
        r"\b(?:NASDAQ|NYSE|NYSEAMERICAN|NYSEARCA)\s*[:：]\s*"
        r"(?P<code>[A-Z][A-Z0-9.-]{0,14})\b",
        re.I,
    ),
)


def _setting(settings: object, name: str, default: object) -> object:
    if isinstance(settings, Mapping):
        return settings.get(name, default)
    return getattr(settings, name, default)


def _utc_text(value: datetime) -> str:
    if value.tzinfo is None or value.utcoffset() is None:
        raise ValueError("datetime must be timezone-aware")
    return value.astimezone(UTC).isoformat(timespec="milliseconds").replace(
        "+00:00", "Z"
    )


def _expiry(requested_at: datetime) -> str:
    return _utc_text(requested_at.astimezone(UTC) + timedelta(hours=24))


def _normalized(value: object) -> str:
    try:
        return normalize_alias(str(value or ""))
    except ValueError:
        return ""


def _normalized_name(value: object) -> str:
    return _normalized(str(value or "").strip().rstrip(".,。"))


def _hk_code(value: object) -> str:
    digits = str(value or "").strip().split(".", 1)[0]
    if not digits.isdigit():
        return ""
    return digits.lstrip("0").zfill(4) if int(digits) < 10000 else digits


def _cn_code(value: object) -> str:
    raw = str(value or "").strip()
    if re.fullmatch(r"\d{1,6}(?:\.0+)?", raw):
        raw = raw.split(".", 1)[0]
    return raw.zfill(6) if raw.isdigit() and len(raw) <= 6 else ""


def _date_text(value: object) -> str:
    if value is None:
        return ""
    if hasattr(value, "date") and callable(value.date):
        try:
            return value.date().isoformat()
        except (TypeError, ValueError):
            pass
    if hasattr(value, "isoformat") and callable(value.isoformat):
        try:
            return value.isoformat()
        except (TypeError, ValueError):
            pass
    raw = str(value).strip()
    compact = re.fullmatch(r"(?P<year>\d{4})(?P<month>\d{2})(?P<day>\d{2})", raw)
    if compact:
        return "{year}-{month}-{day}".format(**compact.groupdict())
    return raw[:10]


def _identity_assertion(
    *,
    canonical_symbol: str,
    display_name: str,
    asset_type: str,
    market: str,
    exchange: str,
    currency: str,
    country_code: str,
    aliases: list[str],
    source_key: str,
    source_type: str,
    source_role: str,
    source_url: str,
    requested_at: datetime,
    **extra,
) -> dict:
    return {
        "canonical_symbol": canonical_symbol,
        "display_name": display_name,
        "asset_type": asset_type,
        "market": market,
        "exchange": exchange,
        "currency": currency,
        "country_code": country_code,
        "listing_status": "active",
        "aliases": [item for item in aliases if str(item or "").strip()],
        "source_key": source_key,
        "source_type": source_type,
        "source_role": source_role,
        "source_url": source_url,
        "observed_at": _utc_text(requested_at),
        "expires_at": _expiry(requested_at),
        **extra,
    }


class NasdaqTraderInstrumentDiscoverySource:
    """纳斯达克官方 Symbol Directory，覆盖美国交易所股票和 ETF。"""

    source_key = "nasdaq_trader_symbol_directory"

    def __init__(self, *, http_client=None):
        self.http = http_client or SafeHTTPClient()
        self._lock = threading.RLock()
        self._rows = None

    def _load_rows(self) -> tuple[dict, ...]:
        with self._lock:
            if self._rows is not None:
                return self._rows
            rows = []
            for url, exchange in (
                (NASDAQ_LISTED_URL, "XNAS"),
                (NASDAQ_OTHER_URL, ""),
            ):
                result = self.http.get(
                    url,
                    headers={"User-Agent": "CollectInfo-FinancialDiscovery/1.0"},
                )
                for item in csv.DictReader(io.StringIO(result.text), delimiter="|"):
                    if not isinstance(item, dict):
                        continue
                    symbol = str(
                        item.get("Symbol") or item.get("ACT Symbol") or ""
                    ).strip().upper()
                    if not symbol or symbol.startswith("FILE CREATION TIME"):
                        continue
                    item = dict(item)
                    item["_symbol"] = symbol
                    item["_exchange"] = exchange or _EXCHANGE_CODES.get(
                        str(item.get("Exchange") or "").strip().upper(), ""
                    )
                    item["_source_url"] = url
                    rows.append(item)
            self._rows = tuple(rows)
            return self._rows

    @staticmethod
    def _matches(query: str, item: Mapping[str, object]) -> bool:
        query_value = str(query or "").strip().upper()
        symbol_match = _US_SYMBOL.fullmatch(query_value)
        if symbol_match and symbol_match.group("code").upper() == item["_symbol"]:
            return True
        name = str(item.get("Security Name") or "").strip()
        short_name = re.split(r"\s+-\s+", name, maxsplit=1)[0]
        normalized = _normalized_name(query)
        return bool(normalized) and normalized in {
            _normalized_name(name),
            _normalized_name(short_name),
        }

    def search(self, query: str, *, requested_at: datetime, request_id: str) -> list[dict]:
        del request_id
        if _HK_SYMBOL.fullmatch(str(query or "").strip()):
            return []
        assertions = []
        for item in self._load_rows():
            if not self._matches(query, item):
                continue
            if str(item.get("Test Issue") or "").strip().upper() == "Y":
                continue
            exchange = str(item.get("_exchange") or "")
            name = str(item.get("Security Name") or "").strip()
            if not exchange or (str(item.get("ETF") or "").upper() != "Y" and _UNSUPPORTED_SECURITY.search(name)):
                continue
            symbol = str(item["_symbol"])
            asset_type = (
                "etf" if str(item.get("ETF") or "").strip().upper() == "Y" else "equity"
            )
            assertions.append(
                _identity_assertion(
                    canonical_symbol=f"{symbol}.US",
                    display_name=name,
                    asset_type=asset_type,
                    market="US",
                    exchange=exchange,
                    currency="USD",
                    country_code="US",
                    aliases=[symbol, f"{symbol}.US", name],
                    source_key=self.source_key,
                    source_type="exchange",
                    source_role="authoritative",
                    source_url=str(item["_source_url"]),
                    requested_at=requested_at,
                )
            )
        return assertions


class SSEInstrumentDiscoverySource:
    """上海证券交易所官方 A 股名录。"""

    source_key = "sse_official_a_share_list"

    def __init__(self, *, http_client=None):
        self.http = http_client or SafeHTTPClient()
        self._lock = threading.RLock()
        self._rows = None

    def _load_rows(self) -> tuple[dict, ...]:
        with self._lock:
            if self._rows is not None:
                return self._rows
            rows = []
            for stock_type, board in (("1", "main"), ("8", "star")):
                url = f"{SSE_STOCK_LIST_URL}?" + urlencode(
                    {
                        "STOCK_TYPE": stock_type,
                        "REG_PROVINCE": "",
                        "CSRC_CODE": "",
                        "STOCK_CODE": "",
                        "sqlId": "COMMON_SSE_CP_GPJCTPZ_GPLB_GP_L",
                        "COMPANY_STATUS": "2,4,5,7,8",
                        "type": "inParams",
                        "isPagination": "true",
                        "pageHelp.cacheSize": "1",
                        "pageHelp.beginPage": "1",
                        "pageHelp.pageSize": "10000",
                        "pageHelp.pageNo": "1",
                        "pageHelp.endPage": "1",
                    }
                )
                result = self.http.get(
                    url,
                    headers={
                        "Host": "query.sse.com.cn",
                        "Referer": SSE_STOCK_LIST_PAGE_URL,
                        "User-Agent": "CollectInfo-FinancialDiscovery/1.0",
                    },
                )
                payload = json.loads(result.text)
                result_rows = payload.get("result") if isinstance(payload, dict) else None
                if not isinstance(result_rows, list):
                    raise ValueError("SSE stock list schema changed")
                for item in result_rows:
                    if not isinstance(item, Mapping):
                        continue
                    row = dict(item)
                    row["_board"] = board
                    rows.append(row)
            self._rows = tuple(rows)
            return self._rows

    @staticmethod
    def _matches(query: str, item: Mapping[str, object]) -> bool:
        query_text = str(query or "").strip()
        symbol_match = _CN_SYMBOL.fullmatch(query_text)
        code = _cn_code(item.get("A_STOCK_CODE"))
        if symbol_match:
            return (
                symbol_match.group("venue").upper() == "SH"
                and symbol_match.group("code") == code
            )
        if re.fullmatch(r"\d{6}", query_text):
            return False
        normalized = _normalized(query_text)
        names = (
            item.get("SEC_NAME_CN"),
            item.get("SEC_NAME_FULL"),
            item.get("COMPANY_ABBR"),
            item.get("FULL_NAME"),
        )
        return bool(normalized) and normalized in {
            _normalized(name) for name in names if str(name or "").strip()
        }

    def search(self, query: str, *, requested_at: datetime, request_id: str) -> list[dict]:
        del request_id
        query_text = str(query or "").strip()
        explicit = _CN_SYMBOL.fullmatch(query_text)
        if explicit and explicit.group("venue").upper() != "SH":
            return []
        if query_text.upper().endswith((".HK", ".HKG", ".US")):
            return []
        assertions = []
        for item in self._load_rows():
            if not self._matches(query_text, item):
                continue
            code = _cn_code(item.get("A_STOCK_CODE"))
            name = str(item.get("SEC_NAME_CN") or "").strip()
            if not code or not name:
                continue
            aliases = [
                code,
                f"{code}.SH",
                name,
                str(item.get("SEC_NAME_FULL") or "").strip(),
                str(item.get("COMPANY_ABBR") or "").strip(),
                str(item.get("FULL_NAME") or "").strip(),
            ]
            assertions.append(
                _identity_assertion(
                    canonical_symbol=f"{code}.SH",
                    display_name=name,
                    asset_type="equity",
                    market="CN",
                    exchange="XSHG",
                    currency="CNY",
                    country_code="CN",
                    aliases=aliases,
                    source_key=self.source_key,
                    source_type="exchange",
                    source_role="authoritative",
                    source_url=SSE_STOCK_LIST_PAGE_URL,
                    requested_at=requested_at,
                    listed_at=_date_text(item.get("LIST_DATE")),
                    metadata={"security_board": str(item.get("_board") or "")},
                )
            )
        return assertions


class SZSEInstrumentDiscoverySource:
    """深圳证券交易所官方 A 股名录。"""

    source_key = "szse_official_a_share_list"

    def __init__(self, *, http_client=None):
        self.http = http_client or SafeHTTPClient()
        self._lock = threading.RLock()
        self._rows = None

    def _load_rows(self) -> tuple[dict, ...]:
        with self._lock:
            if self._rows is not None:
                return self._rows
            from openpyxl import load_workbook

            url = f"{SZSE_STOCK_LIST_URL}?" + urlencode(
                {"SHOWTYPE": "xlsx", "CATALOGID": "1110", "TABKEY": "tab1"}
            )
            result = self.http.get(
                url,
                headers={
                    "Referer": SZSE_STOCK_LIST_PAGE_URL,
                    "User-Agent": "CollectInfo-FinancialDiscovery/1.0",
                },
            )
            workbook = load_workbook(
                io.BytesIO(result.content), read_only=True, data_only=True
            )
            sheet = workbook.active
            sheet.reset_dimensions()
            headers = None
            rows = []
            for values in sheet.iter_rows(values_only=True):
                labels = [str(item or "").strip() for item in values]
                if headers is None:
                    if "A股代码" in labels and "A股简称" in labels:
                        headers = labels
                    continue
                item = dict(zip(headers, values))
                if _cn_code(item.get("A股代码")):
                    rows.append(item)
            workbook.close()
            if headers is None:
                raise ValueError("SZSE stock list workbook header not found")
            self._rows = tuple(rows)
            return self._rows

    @staticmethod
    def _matches(query: str, item: Mapping[str, object]) -> bool:
        query_text = str(query or "").strip()
        symbol_match = _CN_SYMBOL.fullmatch(query_text)
        code = _cn_code(item.get("A股代码"))
        if symbol_match:
            return (
                symbol_match.group("venue").upper() == "SZ"
                and symbol_match.group("code") == code
            )
        if re.fullmatch(r"\d{6}", query_text):
            return False
        normalized = _normalized(query_text)
        return bool(normalized) and normalized == _normalized(item.get("A股简称"))

    def search(self, query: str, *, requested_at: datetime, request_id: str) -> list[dict]:
        del request_id
        query_text = str(query or "").strip()
        explicit = _CN_SYMBOL.fullmatch(query_text)
        if explicit and explicit.group("venue").upper() != "SZ":
            return []
        if query_text.upper().endswith((".HK", ".HKG", ".US")):
            return []
        assertions = []
        for item in self._load_rows():
            if not self._matches(query_text, item):
                continue
            code = _cn_code(item.get("A股代码"))
            name = str(item.get("A股简称") or "").strip()
            if not code or not name:
                continue
            assertions.append(
                _identity_assertion(
                    canonical_symbol=f"{code}.SZ",
                    display_name=name,
                    asset_type="equity",
                    market="CN",
                    exchange="XSHE",
                    currency="CNY",
                    country_code="CN",
                    aliases=[code, f"{code}.SZ", name],
                    source_key=self.source_key,
                    source_type="exchange",
                    source_role="authoritative",
                    source_url=SZSE_STOCK_LIST_PAGE_URL,
                    requested_at=requested_at,
                    listed_at=_date_text(item.get("A股上市日期")),
                    metadata={
                        "security_board": str(item.get("板块") or "").strip(),
                        "industry": str(item.get("所属行业") or "").strip(),
                    },
                )
            )
        return assertions


class HKEXInstrumentDiscoverySource:
    """港交所官方完整证券名录，覆盖股票、ETF、REIT 等上市产品。"""

    source_key = "hkex_full_list_of_securities"

    def __init__(self, *, http_client=None):
        self.http = http_client or SafeHTTPClient()
        self._lock = threading.RLock()
        self._rows = None

    def _load_rows(self) -> tuple[dict, ...]:
        with self._lock:
            if self._rows is not None:
                return self._rows
            from openpyxl import load_workbook

            def read_rows(url, code_header, name_header):
                result = self.http.get(
                    url,
                    headers={"User-Agent": "CollectInfo-FinancialDiscovery/1.0"},
                )
                workbook = load_workbook(
                    io.BytesIO(result.content), read_only=True, data_only=True
                )
                sheet = workbook.active
                sheet.reset_dimensions()
                headers = None
                parsed = []
                for values in sheet.iter_rows(values_only=True):
                    labels = [str(item or "").strip() for item in values]
                    if headers is None:
                        if code_header in labels and name_header in labels:
                            headers = labels
                        continue
                    item = dict(zip(headers, values))
                    if str(item.get(code_header) or "").strip():
                        parsed.append(item)
                workbook.close()
                if headers is None:
                    raise ValueError(f"HKEX securities workbook header not found: {url}")
                return parsed

            rows = read_rows(HKEX_SECURITIES_URL, "Stock Code", "Name of Securities")
            try:
                chinese_rows = read_rows(HKEX_SECURITIES_ZH_URL, "股份代號", "股份名稱")
            except Exception:
                chinese_rows = []
            chinese_by_code = {
                str(item.get("股份代號") or "").strip().zfill(5): item
                for item in chinese_rows
            }
            for item in rows:
                code = str(item.get("Stock Code") or "").strip().zfill(5)
                chinese = chinese_by_code.get(code) or {}
                item["Chinese Name of Securities"] = str(
                    chinese.get("股份名稱") or ""
                ).strip()
            self._rows = tuple(rows)
            return self._rows

    @staticmethod
    def _asset_type(item: Mapping[str, object]) -> str:
        category = str(item.get("Category") or "").casefold()
        subcategory = str(item.get("Sub-Category") or "").casefold()
        if "exchange traded fund" in subcategory:
            return "etf"
        if "real estate investment trust" in subcategory or "reit" in subcategory:
            return "fund"
        if "equity" in category:
            return "equity"
        return ""

    def search(self, query: str, *, requested_at: datetime, request_id: str) -> list[dict]:
        del request_id
        query_text = str(query or "").strip()
        if query_text.upper().endswith(".US"):
            return []
        match = _HK_SYMBOL.fullmatch(query_text)
        requested_code = _hk_code(match.group("code")) if match else ""
        normalized_query = _normalized(query_text)
        assertions = []
        for item in self._load_rows():
            raw_code = str(item.get("Stock Code") or "").strip().zfill(5)
            code = _hk_code(raw_code)
            name = str(item.get("Name of Securities") or "").strip()
            chinese_name = str(item.get("Chinese Name of Securities") or "").strip()
            official_names = {
                key: value
                for key, value in {
                    "en_short": name,
                    "zh_short": chinese_name,
                    "en_full": str(item.get("Full Name of Issuer") or "").strip(),
                    "zh_full": str(item.get("發行人全名") or "").strip(),
                }.items()
                if value
            }
            former_names = [
                str(item.get(key) or "").strip()
                for key in (
                    "Former Name",
                    "Former Name(s)",
                    "Previous Name",
                    "Previous Name(s)",
                    "舊名",
                    "前稱",
                )
                if str(item.get(key) or "").strip()
            ]
            names = [*official_names.values(), *former_names]
            if requested_code:
                matched = requested_code == code
            else:
                matched = normalized_query in {
                    _normalized(value) for value in [raw_code, code, *names] if value
                }
            asset_type = self._asset_type(item)
            if not matched or not asset_type:
                continue
            currency = str(item.get("Trading Currency") or "HKD").strip().upper()
            assertions.append(
                _identity_assertion(
                    canonical_symbol=f"{code}.HK",
                    display_name=name,
                    asset_type=asset_type,
                    market="XHKG",
                    exchange="XHKG",
                    currency=currency,
                    country_code="HK",
                    aliases=list(dict.fromkeys([raw_code, code, f"{code}.HK", *names])),
                    source_key=self.source_key,
                    source_type="exchange",
                    source_role="authoritative",
                    source_url=HKEX_SECURITIES_URL,
                    requested_at=requested_at,
                    metadata={
                        "isin": str(item.get("ISIN") or "").strip(),
                        "category": str(item.get("Category") or "").strip(),
                        "subcategory": str(item.get("Sub-Category") or "").strip(),
                        "official_names": official_names,
                        "former_names": former_names,
                        "official_name_sources": [
                            HKEX_SECURITIES_URL,
                            HKEX_SECURITIES_ZH_URL,
                        ],
                    },
                )
            )
        return assertions


class AlphaVantageInstrumentDiscoverySource:
    """已授权 Alpha Vantage SYMBOL_SEARCH，仅提供独立身份和 Provider 映射。"""

    source_key = "alpha_vantage_symbol_search"
    canonical_requery = True
    provider_mapping_verifier = True

    def __init__(self, *, settings, http_client=None):
        self.settings = settings
        self.http = http_client or SafeHTTPClient()

    def _authorized_key(self, requested_at: datetime) -> str:
        require_financial_capability("alpha_vantage", self.settings)
        require_provider_authorization(
            "alpha_vantage",
            self.settings,
            at=requested_at,
            provider_profile=_ALPHA_PROFILE,
        )
        key = str(_setting(self.settings, "ALPHA_VANTAGE_API_KEY", "") or "").strip()
        if not key:
            raise PermissionError("alpha_vantage_api_key_missing")
        return key

    @staticmethod
    def _identity(item: Mapping[str, object]) -> tuple[str, str, str, str, str]:
        raw_symbol = str(item.get("1. symbol") or "").strip().upper()
        region = str(item.get("4. region") or "").strip().casefold()
        currency = str(item.get("8. currency") or "").strip().upper()
        type_name = str(item.get("3. type") or "").strip().casefold()
        if raw_symbol.endswith(".HKG") or region == "hong kong":
            code = _hk_code(raw_symbol.rsplit(".", 1)[0])
            canonical = f"{code}.HK" if code else ""
            return canonical, "XHKG", "XHKG", currency or "HKD", "HK"
        if region in {"united states", "usa", "us"}:
            base = raw_symbol.removesuffix(".US")
            canonical = f"{base}.US" if _US_SYMBOL.fullmatch(base) else ""
            return canonical, "US", "", currency or "USD", "US"
        return "", "", "", "", ""

    @staticmethod
    def _asset_type(value: object) -> str:
        normalized = str(value or "").strip().casefold()
        if "etf" in normalized:
            return "etf"
        if "fund" in normalized:
            return "fund"
        if normalized in {"equity", "stock"}:
            return "equity"
        return ""

    @staticmethod
    def _matches(query: str, item: Mapping[str, object], canonical: str) -> bool:
        query_normalized = _normalized_name(query)
        raw_symbol = str(item.get("1. symbol") or "").strip()
        name = str(item.get("2. name") or "").strip()
        return bool(query_normalized) and query_normalized in {
            _normalized_name(raw_symbol),
            _normalized_name(canonical),
            _normalized_name(raw_symbol.split(".", 1)[0]),
            _normalized_name(name),
        }

    def search(self, query: str, *, requested_at: datetime, request_id: str) -> list[dict]:
        del request_id
        key = self._authorized_key(requested_at)
        query_text = str(query or "").strip()
        hk_match = _HK_SYMBOL.fullmatch(query_text)
        us_match = _US_SYMBOL.fullmatch(query_text.upper())
        keywords = (
            str(int(hk_match.group("code")))
            if hk_match and hk_match.group("code").isdigit()
            else (
                query_text.upper().removesuffix(".US")
                if us_match and query_text.upper().endswith(".US")
                else query_text
            )
        )
        url = f"{_ALPHA_PROFILE['api_base_url']}?" + urlencode(
            {
                "function": "SYMBOL_SEARCH",
                "keywords": keywords[:120],
                "apikey": key,
                "datatype": "json",
            }
        )
        result = self.http.get(
            url,
            headers={"User-Agent": "CollectInfo-FinancialDiscovery/1.0"},
        )
        payload = json.loads(result.text)
        if not isinstance(payload, dict):
            raise ValueError("Alpha Vantage symbol search schema changed")
        if payload.get("Note") or payload.get("Information") or payload.get("Error Message"):
            raise RuntimeError("Alpha Vantage symbol search unavailable")
        assertions = []
        for item in payload.get("bestMatches") or []:
            if not isinstance(item, Mapping):
                continue
            canonical, market, exchange, currency, country = self._identity(item)
            asset_type = self._asset_type(item.get("3. type"))
            if not canonical or not asset_type or not self._matches(query, item, canonical):
                continue
            raw_symbol = str(item.get("1. symbol") or "").strip().upper()
            name = str(item.get("2. name") or "").strip()
            base = _identity_assertion(
                canonical_symbol=canonical,
                display_name=name,
                asset_type=asset_type,
                market=market,
                exchange=exchange,
                currency=currency,
                country_code=country,
                aliases=[raw_symbol, canonical, raw_symbol.split(".", 1)[0], name],
                source_key=self.source_key,
                source_type="provider",
                source_role="corroborating",
                source_url=str(_ALPHA_PROFILE["documentation_url"]),
                requested_at=requested_at,
            )
            assertions.append(base)
            assertions.append(
                {
                    **base,
                    "source_key": f"{self.source_key}_mapping",
                    "source_role": "provider",
                    "approved": True,
                    "provider_key": "alpha_vantage",
                    "provider_symbol": raw_symbol,
                }
            )
        return assertions


class YahooInstrumentDiscoverySource:
    """通过 Yahoo 精确搜索补充研究级身份核验和查询映射。"""

    source_key = "yahoo_finance_symbol_search"
    canonical_requery = True
    provider_mapping_verifier = True
    _US_EXCHANGES = {
        "NMS": "XNAS",
        "NGM": "XNAS",
        "NCM": "XNAS",
        "NYQ": "XNYS",
        "ASE": "XASE",
        "PCX": "ARCX",
    }

    def __init__(self, *, settings, sdk=None):
        self.settings = settings
        self._sdk = sdk

    def _authorized_sdk(self, requested_at: datetime):
        require_financial_capability("yahoo", self.settings)
        require_provider_authorization(
            "yahoo",
            self.settings,
            at=requested_at,
            provider_profile=_YAHOO_PROFILE,
        )
        if self._sdk is None:
            import yfinance as sdk

            self._sdk = sdk
        return self._sdk

    @classmethod
    def _identity(cls, item: Mapping[str, object]):
        raw_symbol = str(item.get("symbol") or "").strip().upper()
        quote_type = str(item.get("quoteType") or "").strip().upper()
        asset_type = {"EQUITY": "equity", "ETF": "etf"}.get(quote_type, "")
        if not raw_symbol or not asset_type:
            return None
        if raw_symbol.endswith(".HK"):
            code = _hk_code(raw_symbol.rsplit(".", 1)[0])
            canonical = f"{code}.HK" if code else ""
            return canonical, "XHKG", "XHKG", "HKD", "HK", asset_type
        exchange = cls._US_EXCHANGES.get(
            str(item.get("exchange") or "").strip().upper(), ""
        )
        if exchange and _US_SYMBOL.fullmatch(raw_symbol):
            return f"{raw_symbol}.US", "US", exchange, "USD", "US", asset_type
        return None

    @staticmethod
    def _matches(query: str, item: Mapping[str, object], canonical: str) -> bool:
        raw_symbol = str(item.get("symbol") or "").strip()
        names = (
            item.get("shortname"),
            item.get("longname"),
            item.get("displayName"),
        )
        normalized_query = _normalized_name(query)
        candidates = {
            _normalized_name(raw_symbol),
            _normalized_name(raw_symbol.split(".", 1)[0]),
            _normalized_name(canonical),
            *(_normalized_name(value) for value in names if value),
        }
        return bool(normalized_query) and normalized_query in candidates

    def search(self, query: str, *, requested_at: datetime, request_id: str) -> list[dict]:
        del request_id
        sdk = self._authorized_sdk(requested_at)
        query_text = str(query or "").strip()
        us_match = _US_SYMBOL.fullmatch(query_text.upper())
        keywords = (
            query_text.upper().removesuffix(".US")
            if us_match and query_text.upper().endswith(".US")
            else query_text
        )
        search = sdk.Search(
            keywords,
            max_results=8,
            news_count=0,
            lists_count=0,
            include_cb=False,
            include_nav_links=False,
            recommended=0,
            timeout=int(_setting(self.settings, "FINANCIAL_PROVIDER_TIMEOUT_SECONDS", 20)),
        )
        assertions = []
        for item in getattr(search, "quotes", ()) or ():
            if not isinstance(item, Mapping):
                continue
            identity = self._identity(item)
            if identity is None:
                continue
            canonical, market, exchange, currency, country, asset_type = identity
            if not canonical or not self._matches(query, item, canonical):
                continue
            raw_symbol = str(item.get("symbol") or "").strip().upper()
            name = str(
                item.get("longname") or item.get("shortname") or raw_symbol
            ).strip()
            base = _identity_assertion(
                canonical_symbol=canonical,
                display_name=name,
                asset_type=asset_type,
                market=market,
                exchange=exchange,
                currency=currency,
                country_code=country,
                aliases=[raw_symbol, canonical, raw_symbol.split(".", 1)[0], name],
                source_key=self.source_key,
                source_type="provider",
                source_role="corroborating",
                source_url=str(_YAHOO_PROFILE["documentation_url"]),
                requested_at=requested_at,
            )
            assertions.extend(
                (
                    base,
                    {
                        **base,
                        "source_key": f"{self.source_key}_mapping",
                        "source_role": "provider",
                        "approved": True,
                        "provider_key": "yahoo",
                        "provider_symbol": raw_symbol,
                    },
                )
            )
        return assertions


class AKShareAInstrumentDiscoverySource:
    """AKShare 已授权上游的沪深 A 股单标的身份与 Provider 映射。"""

    source_key = "akshare_eastmoney_a_share_quote_identity"
    canonical_requery = True

    def __init__(self, *, settings, http_client=None):
        self.settings = settings
        self.http = http_client or SafeHTTPClient()

    def _authorized(self, requested_at: datetime) -> None:
        require_financial_capability("akshare_cn", self.settings)
        require_provider_authorization(
            "akshare_cn",
            self.settings,
            at=requested_at,
            provider_profile=_AKSHARE_PROFILE,
        )

    def search(self, query: str, *, requested_at: datetime, request_id: str) -> list[dict]:
        del request_id
        match = _CN_SYMBOL.fullmatch(str(query or "").strip())
        if match is None:
            return []
        self._authorized(requested_at)
        code = match.group("code")
        venue = match.group("venue").upper()
        market_marker = "1" if venue == "SH" else "0"
        exchange = "XSHG" if venue == "SH" else "XSHE"
        quote_query = urlencode(
            {
                "secid": f"{market_marker}.{code}",
                "fields": "f57,f58,f107,f152",
                "fltt": "2",
                "invt": "2",
            }
        )
        errors = []
        quote_data = None
        for endpoint in (
            EASTMONEY_HK_QUOTE_IDENTITY_URL,
            EASTMONEY_HK_QUOTE_IDENTITY_FALLBACK_URL,
        ):
            try:
                result = self.http.get(
                    f"{endpoint}?{quote_query}",
                    headers={"User-Agent": "CollectInfo-FinancialDiscovery/1.0"},
                )
                payload = json.loads(result.text)
            except Exception as error:
                errors.append(error)
                continue
            candidate = (
                payload.get("data")
                if isinstance(payload, dict)
                and isinstance(payload.get("data"), Mapping)
                else None
            )
            if (
                candidate
                and _cn_code(candidate.get("f57")) == code
                and str(
                    candidate.get("f107")
                    if candidate.get("f107") is not None
                    else ""
                )
                == market_marker
                and str(candidate.get("f58") or "").strip()
            ):
                quote_data = candidate
                break
        if quote_data is None:
            if len(errors) == 2:
                raise errors[-1]
            return []

        name = str(quote_data.get("f58") or "").strip()
        canonical = f"{code}.{venue}"
        base = _identity_assertion(
            canonical_symbol=canonical,
            display_name=name,
            # Quote identity does not authoritatively classify the security.
            # The exchange assertion must supply the asset type at admission.
            asset_type="",
            market="CN",
            exchange=exchange,
            currency="CNY",
            country_code="CN",
            aliases=[code, canonical, name],
            source_key=self.source_key,
            source_type="provider",
            source_role="corroborating",
            source_url=f"https://quote.eastmoney.com/{venue.casefold()}{code}.html",
            requested_at=requested_at,
        )
        return [
            base,
            {
                **base,
                "source_key": f"{self.source_key}_mapping",
                "source_role": "provider",
                "approved": True,
                "provider_key": "akshare_cn",
                "provider_symbol": code,
            },
        ]


class AKShareHKInstrumentDiscoverySource:
    """AKShare 已授权上游的港股单标的证券资料，不调用全量行情列表。"""

    source_key = "akshare_eastmoney_hk_security_profile"
    canonical_requery = True

    def __init__(self, *, settings, http_client=None):
        self.settings = settings
        self.http = http_client or SafeHTTPClient()

    def _authorized(self, requested_at: datetime) -> None:
        require_financial_capability("akshare_cn", self.settings)
        require_provider_authorization(
            "akshare_cn",
            self.settings,
            at=requested_at,
            provider_profile=_AKSHARE_PROFILE,
        )

    @staticmethod
    def _asset_type(value: object) -> str:
        normalized = str(value or "").strip().casefold()
        if "etf" in normalized or "交易所买卖基金" in normalized:
            return "etf"
        if "reit" in normalized or "房地产投资信托" in normalized:
            return "fund"
        return "equity" if normalized else ""

    def search(self, query: str, *, requested_at: datetime, request_id: str) -> list[dict]:
        del request_id
        match = _HK_SYMBOL.fullmatch(str(query or "").strip())
        if match is None:
            return []
        self._authorized(requested_at)
        raw_code = str(match.group("code") or "").zfill(5)
        canonical_code = _hk_code(raw_code)
        url = f"{EASTMONEY_HK_PROFILE_URL}?" + urlencode(
            {
                "reportName": "RPT_HKF10_INFO_SECURITYINFO",
                "columns": (
                    "SECUCODE,SECURITY_CODE,SECURITY_NAME_ABBR,SECURITY_TYPE,"
                    "LISTING_DATE,ISIN_CODE,BOARD,TRADE_MARKET"
                ),
                "quoteColumns": "",
                "filter": f'(SECUCODE="{raw_code}.HK")',
                "pageNumber": "1",
                "pageSize": "1",
                "sortTypes": "",
                "sortColumns": "",
                "source": "F10",
                "client": "PC",
            }
        )
        result = self.http.get(
            url,
            headers={"User-Agent": "CollectInfo-FinancialDiscovery/1.0"},
        )
        payload = json.loads(result.text)
        rows = ((payload.get("result") or {}).get("data") or []) if isinstance(payload, dict) else []
        if not rows:
            quote_query = urlencode(
                {
                    "secid": f"116.{raw_code}",
                    "fields": "f57,f58,f107,f152",
                    "fltt": "2",
                    "invt": "2",
                }
            )
            quote_errors = []
            for endpoint in (
                EASTMONEY_HK_QUOTE_IDENTITY_URL,
                EASTMONEY_HK_QUOTE_IDENTITY_FALLBACK_URL,
            ):
                try:
                    quote_result = self.http.get(
                        f"{endpoint}?{quote_query}",
                        headers={
                            "User-Agent": "CollectInfo-FinancialDiscovery/1.0"
                        },
                    )
                    quote_payload = json.loads(quote_result.text)
                except Exception as error:
                    quote_errors.append(error)
                    continue
                quote_data = (
                    quote_payload.get("data")
                    if isinstance(quote_payload, dict)
                    and isinstance(quote_payload.get("data"), Mapping)
                    else None
                )
                if not quote_data or str(quote_data.get("f107") or "") != "116":
                    continue
                rows = [
                    {
                        "SECURITY_CODE": quote_data.get("f57"),
                        "SECURITY_NAME_ABBR": quote_data.get("f58"),
                        # The quote identity endpoint does not authoritatively
                        # classify ETF/fund/equity.  Leave the field unknown;
                        # the exchange assertion must supply it at admission.
                        "SECURITY_TYPE": "",
                    }
                ]
                break
            if not rows and len(quote_errors) == 2:
                raise quote_errors[-1]
        assertions = []
        for item in rows[:1]:
            if not isinstance(item, Mapping):
                continue
            returned_code = str(item.get("SECURITY_CODE") or "").strip().zfill(5)
            asset_type = self._asset_type(item.get("SECURITY_TYPE"))
            name = str(item.get("SECURITY_NAME_ABBR") or "").strip()
            if returned_code != raw_code or not name:
                continue
            base = _identity_assertion(
                canonical_symbol=f"{canonical_code}.HK",
                display_name=name,
                asset_type=asset_type,
                market="XHKG",
                exchange="XHKG",
                currency="HKD",
                country_code="HK",
                aliases=[raw_code, canonical_code, f"{canonical_code}.HK", name],
                source_key=self.source_key,
                source_type="provider",
                source_role="corroborating",
                source_url=(
                    "https://emweb.securities.eastmoney.com/PC_HKF10/pages/home/"
                    f"index.html?code={raw_code}"
                ),
                requested_at=requested_at,
                listed_at=str(item.get("LISTING_DATE") or "")[:10],
                metadata={
                    "isin": str(item.get("ISIN_CODE") or "").strip(),
                    "board": str(item.get("BOARD") or "").strip(),
                },
            )
            assertions.append(base)
            assertions.append(
                {
                    **base,
                    "source_key": f"{self.source_key}_mapping",
                    "source_role": "provider",
                    "approved": True,
                    "provider_key": "akshare_cn",
                    "provider_symbol": raw_code,
                }
            )
        return assertions


class SerpAPIInstrumentHintSource:
    """可选搜索兜底；输出只能作为再次查询结构化来源的候选词。"""

    source_key = "serpapi_instrument_hint"
    fallback_only = True

    def __init__(self, *, client=None):
        self.client = client or SerpAPIClient()

    @staticmethod
    def _suggestions(text: str) -> list[str]:
        suggestions = []
        for pattern in _SUGGESTION_PATTERNS:
            for match in pattern.finditer(str(text or "")):
                code = str(match.group("code") or "").upper()
                suggestion = (
                    f"{_hk_code(code)}.HK" if code.isdigit() else f"{code}.US"
                )
                if suggestion and suggestion not in suggestions:
                    suggestions.append(suggestion)
        return suggestions[:3]

    def search(self, query: str, *, requested_at: datetime, request_id: str) -> list[dict]:
        del request_id
        results = self.client.search(
            f"{str(query or '')[:120]} 股票 基金 ticker exchange official",
            recency_days=0,
        )
        hints = []
        for item in results[:10]:
            title = str(item.get("title") or "").strip()
            summary = str(item.get("summary") or "").strip()
            url = str(item.get("url") or "").strip()
            suggestions = self._suggestions(f"{title} {summary} {url}")
            hints.append(
                {
                    "source_key": self.source_key,
                    "source_type": "search_engine",
                    "source_role": "web_hint",
                    "source_url": url,
                    "title": title,
                    "summary": summary,
                    "suggested_queries": suggestions,
                    "observed_at": _utc_text(requested_at),
                    "expires_at": _expiry(requested_at),
                }
            )
        return hints


def default_external_instrument_sources(*, settings) -> tuple[object, ...]:
    sources: list[object] = [
        NasdaqTraderInstrumentDiscoverySource(),
        SSEInstrumentDiscoverySource(),
        SZSEInstrumentDiscoverySource(),
        HKEXInstrumentDiscoverySource(),
        AKShareAInstrumentDiscoverySource(settings=settings),
        AKShareHKInstrumentDiscoverySource(settings=settings),
        AlphaVantageInstrumentDiscoverySource(settings=settings),
        YahooInstrumentDiscoverySource(settings=settings),
    ]
    if bool(
        _setting(settings, "FINANCIAL_INSTRUMENT_SEARCH_FALLBACK_ENABLED", True)
    ):
        sources.append(SerpAPIInstrumentHintSource())
    return tuple(sources)


__all__ = [
    "NASDAQ_LISTED_URL",
    "NASDAQ_OTHER_URL",
    "HKEX_SECURITIES_URL",
    "HKEX_SECURITIES_ZH_URL",
    "SSE_STOCK_LIST_URL",
    "SSE_STOCK_LIST_PAGE_URL",
    "SZSE_STOCK_LIST_URL",
    "SZSE_STOCK_LIST_PAGE_URL",
    "EASTMONEY_HK_PROFILE_URL",
    "EASTMONEY_HK_QUOTE_IDENTITY_URL",
    "EASTMONEY_HK_QUOTE_IDENTITY_FALLBACK_URL",
    "NasdaqTraderInstrumentDiscoverySource",
    "SSEInstrumentDiscoverySource",
    "SZSEInstrumentDiscoverySource",
    "HKEXInstrumentDiscoverySource",
    "AlphaVantageInstrumentDiscoverySource",
    "YahooInstrumentDiscoverySource",
    "AKShareAInstrumentDiscoverySource",
    "AKShareHKInstrumentDiscoverySource",
    "SerpAPIInstrumentHintSource",
    "default_external_instrument_sources",
]
